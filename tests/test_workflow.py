from pathlib import Path
import sqlite3

from app import create_app
from app.services.email_sender import resolve_smtp_settings
from app.services.employee_importer import build_employee_record, normalize_email, validate_employee_row
from app.services.signature_generator import ICON_BLUE, _font, generate_signature_image
from app.services.text_formatting import format_job_title, format_person_name
from app.services.admin_auth import authenticate_user
from app.services.request_log import append_request_log, successful_request_emails
from app.services.workbook_store import (
    find_employee_by_email,
    list_employees,
    merge_emails_from_workbook,
    replace_employee_workbook_from_upload,
    save_employee_to_workbook,
)


def _corporate_database(path, *, published=True):
    connection = sqlite3.connect(path)
    connection.executescript("""
        CREATE TABLE people (
            id INTEGER PRIMARY KEY, nome TEXT, cargo TEXT, email TEXT, celular TEXT,
            ativo TEXT, ativo_rh INTEGER, email_valido INTEGER, status_validacao TEXT,
            mat TEXT, razao_social TEXT, sharepoint_upn TEXT, sharepoint_ativo INTEGER
        );
        CREATE TABLE rh_importacoes (
            id INTEGER PRIMARY KEY, source_hash TEXT, arquivo_origem TEXT,
            imported_at TEXT, imported_by TEXT, status TEXT, ativos INTEGER,
            com_email INTEGER, sem_email INTEGER
        );
        CREATE VIEW rh_assinaturas_colaboradores AS SELECT * FROM people;
    """)
    connection.executemany(
        "INSERT INTO people VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        [
            (1, "Maria Teste", "Analista", "maria@empresa.com", "1111", "SIM", 1, 1, "APTO", "1", "Empresa", "maria@lorac.test", 1),
            (2, "Pessoa Inativa", "Gestor", "inativa@empresa.com", "", "NÃO", 0, 1, "INATIVO", "2", "Empresa", "", 0),
            (3, "Sem E-mail Válido", "Operador", "pendente@empresa.com", "", "SIM", 1, 0, "SEM_EMAIL", "3", "Empresa", "pendente@lorac.test", 1),
        ],
    )
    if published:
        connection.execute(
            "INSERT INTO rh_importacoes VALUES (1, 'abc', 'rh.xlsx', '2026-09-30T12:00:00Z', 'ti', 'PUBLICADO', 2, 1, 1)"
        )
    connection.commit()
    connection.close()
    return path


def test_normalize_email():
    assert normalize_email("  Ana.Maria@Empresa.COM ") == "ana.maria@empresa.com"


def test_formats_uppercase_employee_name_and_job_title():
    assert format_person_name("ROBERSON AUGUSTO DE SOUZA") == "Roberson Augusto de Souza"
    assert format_job_title("COORDENADOR DE T.I.") == "Coordenador de T.I"
    assert format_job_title("COORDENADOR DE T.I") == "Coordenador de T.I"
    assert format_job_title("COORDENADOR DE TI.") == "Coordenador de TI"
    assert format_job_title("COORDENADOR DE TI") == "Coordenador de TI"
    assert format_job_title("ANALISTA DE RH") == "Analista de RH"


def test_validate_employee_row_rejects_invalid_email():
    row = {"nome": "João Silva", "cargo": "Analista", "email": "joao@invalid", "ativo": "SIM"}
    issues = validate_employee_row(row)
    assert "email" in issues


def test_validate_employee_row_allows_missing_email_during_import():
    row = {"nome": "João Silva", "cargo": "Analista", "email": "", "ativo": "SIM"}

    assert validate_employee_row(row, require_email=False) == {}


def test_validate_employee_row_still_rejects_filled_invalid_email_during_import():
    row = {"nome": "João Silva", "cargo": "Analista", "email": "joao@invalid", "ativo": "SIM"}

    assert "email" in validate_employee_row(row, require_email=False)


def test_build_employee_record_accepts_optional_celular():
    record = build_employee_record({
        "nome": "Ana Silva",
        "cargo": "Analista",
        "email": "ana@empresa.com",
        "celular": "+55 41 99999-0000",
    })
    assert record["phone"] == "+55 41 99999-0000"


def test_generate_signature_image_creates_png(tmp_path):
    target = tmp_path / "assinatura.png"
    result = generate_signature_image(
        name="Maria da Silva",
        role="Analista de Negócios",
        email="maria.silva@empresa.com",
        output_path=target,
    )
    assert result.exists()
    assert result.suffix.lower() == ".png"
    assert result.stat().st_size > 0
    from PIL import Image

    with Image.open(result) as generated:
        assert generated.size == (532, 173)
        # The approved PPT no longer contains the social-media icon strip.
        assert generated.crop((400, 145, 500, 173)).getcolors(maxcolors=1) == [(2800, (255, 255, 255))]
        # Its separator extends through the right side of the signature.
        assert generated.getpixel((500, 69)) == ICON_BLUE


def test_signature_uses_bundled_poppins_bold():
    assert _font(12).getname() == ("Poppins", "Bold")


def test_public_lookup_route(tmp_path):
    database_path = _corporate_database(tmp_path / "corporate.db")
    app = create_app(test_config={"TESTING": True, "SHARED_SQLITE_PATH": str(database_path)})
    with app.test_client() as client:
        response = client.get("/health")
        assert response.status_code == 200
        assert response.json["status"] == "ok"
        home = client.get("/")
        assert 'href="/admin/login"' in home.text
        assert "Área administrativa" in home.text
        assert "entre em contato com o RH" in home.text
        assert "Opções &gt; Email &gt; Assinaturas" in home.text
        assert "respostas e encaminhamentos" in home.text


def test_public_lookup_accepts_only_active_employee_with_valid_email(tmp_path):
    database_path = _corporate_database(tmp_path / "corporate.db")
    app = create_app(test_config={"TESTING": True, "SHARED_SQLITE_PATH": str(database_path)})

    with app.test_client() as client:
        active = client.post("/consultar", data={"email": "MARIA@EMPRESA.COM"})
        inactive = client.post("/consultar", data={"email": "inativa@empresa.com"})
        invalid = client.post("/consultar", data={"email": "pendente@empresa.com"})

    assert "Maria Teste" in active.text
    assert "Nenhum cadastro ativo" in inactive.text
    assert "ainda não possui um e-mail válido" in invalid.text


def test_missing_publication_shows_exact_equalization_warning(tmp_path):
    database_path = _corporate_database(tmp_path / "corporate.db", published=False)
    app = create_app(test_config={"TESTING": True, "SHARED_SQLITE_PATH": str(database_path)})

    with app.test_client() as client:
        response = client.post("/consultar", data={"email": "maria@empresa.com"})

    assert response.status_code == 503
    assert (
        "A base de dados ainda não está equalizada. Solicite à TI que execute a rotina no projeto de "
        "equalização de usuários, sistemas e equipamentos."
    ) in response.text


def test_duplicate_active_email_is_blocked_for_every_employee(tmp_path):
    database_path = _corporate_database(tmp_path / "corporate.db")
    with sqlite3.connect(database_path) as connection:
        connection.execute(
            "INSERT INTO people VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (4, "Outra Pessoa", "Gestor", "maria@empresa.com", "", "SIM", 1, 1, "APTO", "4", "Empresa", "outra@lorac.test", 1),
        )
    app = create_app(test_config={"TESTING": True, "SHARED_SQLITE_PATH": str(database_path)})

    with app.test_client() as client:
        response = client.post("/consultar", data={"email": "maria@empresa.com"})

    assert "Nenhum cadastro ativo" in response.text


def test_unavailable_database_does_not_fall_back_to_legacy_workbook(tmp_path):
    from openpyxl import Workbook

    workbook_path = tmp_path / "legacy.xlsx"
    workbook = Workbook()
    workbook.active.append(["Nome", "Cargo", "Email", "Ativo"])
    workbook.active.append(["Legado", "Analista", "legado@empresa.com", "SIM"])
    workbook.save(workbook_path)
    app = create_app(test_config={
        "TESTING": True,
        "SHARED_SQLITE_PATH": str(tmp_path / "missing.db"),
        "EMPLOYEE_WORKBOOK_PATH": str(workbook_path),
    })

    with app.test_client() as client:
        response = client.post("/consultar", data={"email": "legado@empresa.com"})

    assert response.status_code == 503
    assert "base corporativa está indisponível" in response.text
    assert "Legado" not in response.text


def test_send_uses_email_returned_by_sqlite(tmp_path, monkeypatch):
    database_path = _corporate_database(tmp_path / "corporate.db")
    sent_to = []
    monkeypatch.setattr(
        "app.routes.send_signature_email",
        lambda email, *_args, **_kwargs: sent_to.append(email) or {"status": "simulated"},
    )
    app = create_app(test_config={
        "TESTING": True,
        "SECRET_KEY": "test-secret",
        "SHARED_SQLITE_PATH": str(database_path),
        "GENERATED_FILES_PATH": str(tmp_path / "generated"),
        "REQUEST_LOG_PATH": str(tmp_path / "requests.jsonl"),
    })

    with app.test_client() as client:
        confirmation = client.post("/consultar", data={"email": "MARIA@EMPRESA.COM"})
        token = confirmation.text.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
        client.post("/solicitar", data={
            "csrf_token": token,
            "email": "MARIA@EMPRESA.COM",
            "phone": "1111",
        })

    assert sent_to == ["maria@empresa.com"]


def test_compose_mounts_shared_database_read_only():
    compose = (Path(__file__).resolve().parent.parent / "compose.yaml").read_text(encoding="utf-8")
    assert "/home/administrator/Desktop/codes/ti_dinho_slack/shared_data:/shared:ro" in compose


def test_deploy_script_prepulls_image_and_has_portainer_timeouts():
    root = Path(__file__).resolve().parent.parent
    script = root.joinpath("scripts", "commit-deploy.ps1").read_text(encoding="utf-8")
    config = root.joinpath("deploy.config.psd1").read_text(encoding="utf-8")

    assert "Invoke-PortainerImagePull" in script
    assert "'X-Registry-Auth'" in script
    assert "Update-PortainerContainer" in script
    assert "/rename?name=" in script
    assert "Rollback do container anterior concluido" in script
    assert "TimeoutSec" in script
    assert "RegistryId = 1" in config
    assert "PullTimeoutSeconds = 300" in config
    assert "StackId" not in config
    assert "/api/stacks/" not in script


def test_workbook_is_the_employee_source_of_truth(tmp_path):
    from openpyxl import Workbook

    workbook_path = tmp_path / "colaboradores.xlsx"
    workbook = Workbook()
    workbook.active.append(["Nome", "Cargo", "Email", "Celular", "Ativo"])
    workbook.active.append(["Pessoa Ativa", "Analista", "ativa@empresa.com", "123", "SIM"])
    workbook.active.append(["Pessoa sem E-mail", "Operador", "", "", "SIM"])
    workbook.active.append(["Pessoa Inativa", "Gestor", "inativa@empresa.com", "", "NÃO"])
    workbook.save(workbook_path)

    employees = list_employees(workbook_path)
    assert len(employees) == 3
    assert find_employee_by_email(workbook_path, "ATIVA@EMPRESA.COM")["full_name"] == "Pessoa Ativa"
    assert find_employee_by_email(workbook_path, "inativa@empresa.com") is None


def test_resolve_smtp_settings_supports_sgq_env_names(monkeypatch):
    monkeypatch.delenv("SMTP_HOST", raising=False)
    monkeypatch.delenv("SMTP_PORT", raising=False)
    monkeypatch.delenv("SMTP_USERNAME", raising=False)
    monkeypatch.delenv("SMTP_PASSWORD", raising=False)
    monkeypatch.delenv("SMTP_FROM", raising=False)
    monkeypatch.delenv("SMTP_USE_TLS", raising=False)

    monkeypatch.setenv("SGQ_SMTP_HOST", "smtp.emailnuvem.com.br")
    monkeypatch.setenv("SGQ_SMTP_PORT", "465")
    monkeypatch.setenv("SGQ_SMTP_USER", "suporte.dinho@dinhodistribuidora.com.br")
    monkeypatch.setenv("SGQ_SMTP_PASSWORD", "secret")
    monkeypatch.setenv("SGQ_FROM_EMAIL", "suporte.dinho@dinhodistribuidora.com.br")
    monkeypatch.setenv("SGQ_SMTP_TLS", "false")
    monkeypatch.setenv("SGQ_SMTP_SSL", "true")

    settings = resolve_smtp_settings()

    assert settings["host"] == "smtp.emailnuvem.com.br"
    assert settings["port"] == 465
    assert settings["username"] == "suporte.dinho@dinhodistribuidora.com.br"
    assert settings["password"] == "secret"
    assert settings["from_email"] == "suporte.dinho@dinhodistribuidora.com.br"
    assert settings["use_tls"] is False
    assert settings["use_ssl"] is True


def test_configured_internal_users_authenticate():
    users_path = Path(__file__).resolve().parent.parent / "config" / "users.json"
    assert authenticate_user(users_path, "suporte.dinho", "suporte.dinho")["username"] == "suporte.dinho"
    assert authenticate_user(users_path, "rh", "rh")["username"] == "rh"
    assert authenticate_user(users_path, "rh", "senha-incorreta") is None


def test_save_employee_adds_phone_columns_without_losing_existing_data(tmp_path):
    from openpyxl import Workbook, load_workbook

    workbook_path = tmp_path / "colaboradores.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["MAT", "Nome", "Cargo"])
    sheet.append([10, "Ana Silva", "Analista"])
    workbook.save(workbook_path)

    save_employee_to_workbook(workbook_path, {
        "full_name": "Ana Silva",
        "job_title": "Analista Sênior",
        "email": "ana@dinhodistribuidora.com.br",
        "phone": "+55 41 99999-0000",
        "active": True,
    }, original={"full_name": "Ana Silva", "email": ""})

    saved = load_workbook(workbook_path, data_only=True).active
    assert saved.cell(2, 1).value == 10
    values = {saved.cell(1, column).value: saved.cell(2, column).value for column in range(1, saved.max_column + 1)}
    assert values["Email"] == "ana@dinhodistribuidora.com.br"
    assert values["Celular"] == "+55 41 99999-0000"
    assert values["Cargo"] == "Analista Sênior"


def test_xlsx_import_replaces_entire_base_and_preserves_every_column(tmp_path):
    from io import BytesIO
    from openpyxl import Workbook, load_workbook

    workbook_path = tmp_path / "colaboradores.xlsx"
    old = Workbook()
    old.active.append(["Nome", "Cargo", "Email"])
    old.active.append(["Registro antigo", "Cargo antigo", "antigo@empresa.com"])
    old.save(workbook_path)

    uploaded = Workbook()
    uploaded.active.append([
        "MAT", "Nome", "Razão Social", "Lotação", "Admissão", "CPF",
        "Cargo", "Email", "Celular", "Ativo", "Campo adicional",
    ])
    uploaded.active.append([
        123, "Pessoa Nova", "Empresa", "Unidade", "01/01/2026", "000",
        "Cargo atualizado", "nova@empresa.com", "9999", "SIM", "Preservado",
    ])
    payload = BytesIO()
    uploaded.save(payload)

    result = replace_employee_workbook_from_upload(workbook_path, payload)

    saved = load_workbook(workbook_path, data_only=True).active
    headers = [cell.value for cell in saved[1]]
    values = {headers[index]: saved.cell(2, index + 1).value for index in range(len(headers))}
    assert saved.max_row == 2
    assert values["Nome"] == "Pessoa Nova"
    assert values["Cargo"] == "Cargo atualizado"
    assert values["Campo adicional"] == "Preservado"
    assert "Registro antigo" not in {cell.value for row in saved.iter_rows() for cell in row}
    assert list(tmp_path.joinpath("backups").glob("*.xlsx"))
    assert len(result["sha256"]) == 64
    assert result["size"] == workbook_path.stat().st_size


def test_admin_area_lists_pending_upload_without_changing_published_links(tmp_path):
    from io import BytesIO
    from openpyxl import Workbook

    database_path = _corporate_database(tmp_path / "corporate.db")
    before = database_path.read_bytes()
    users_path = Path(__file__).resolve().parent.parent / "config" / "users.json"
    app = create_app(test_config={
        "TESTING": True,
        "SECRET_KEY": "test-secret",
        "SHARED_SQLITE_PATH": str(database_path),
        "PENDING_IMPORTS_PATH": str(tmp_path / "pending"),
        "ADMIN_USERS_PATH": str(users_path),
        "REQUEST_LOG_PATH": str(tmp_path / "requests.jsonl"),
        "GENERATED_FILES_PATH": str(tmp_path / "generated"),
    })

    with app.test_client() as client:
        assert client.get("/admin").status_code == 302
        login_page = client.get("/admin/login")
        token = login_page.text.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
        login = client.post("/admin/login", data={"username": "rh", "password": "rh", "csrf_token": token})
        assert login.status_code == 302

        page = client.get("/admin/import")
        token = page.text.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
        workbook = Workbook()
        workbook.active.append(["Nome", "Cargo", "Email", "Ativo"])
        workbook.active.append(["Nova Pessoa", "Analista", "nova@empresa.com", "SIM"])
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)
        created = client.post("/admin/import", data={
            "csrf_token": token,
            "file": (payload, "rh_nova.xlsx"),
        }, content_type="multipart/form-data", follow_redirects=True)
        assert "Carga pendente recebida" in created.text
        assert "rh_nova.xlsx" in created.text

    assert database_path.read_bytes() == before
    assert list((tmp_path / "pending").glob("*.xlsx"))


def test_employee_can_update_phone_before_generating_signature(tmp_path, monkeypatch):
    database_path = _corporate_database(tmp_path / "corporate.db")
    before = database_path.read_bytes()
    log_path = tmp_path / "requests.jsonl"
    app = create_app(test_config={
        "TESTING": True,
        "SECRET_KEY": "test-secret",
        "SHARED_SQLITE_PATH": str(database_path),
        "REQUEST_LOG_PATH": str(log_path),
        "GENERATED_FILES_PATH": str(tmp_path / "generated"),
    })
    monkeypatch.setattr("app.routes.send_signature_email", lambda *args, **kwargs: {"status": "simulated"})

    with app.test_client() as client:
        confirmation = client.post("/consultar", data={"email": "maria@empresa.com"})
        assert confirmation.status_code == 200
        assert 'name="phone"' in confirmation.text
        assert 'value="1111"' in confirmation.text
        assert "Caso seja necessária a edição de cargo" in confirmation.text
        token = confirmation.text.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
        sent = client.post("/solicitar", data={
            "csrf_token": token,
            "email": "maria@empresa.com",
            "phone": "+55 41 99999-0000",
        })
        assert sent.status_code == 200
        assert "maria@empresa.com" in successful_request_emails(log_path)

    assert database_path.read_bytes() == before


def test_admin_dashboard_marks_generated_signature_and_has_sortable_headers(tmp_path):
    database_path = _corporate_database(tmp_path / "corporate.db")
    log_path = tmp_path / "requests.jsonl"
    append_request_log(log_path, "maria@empresa.com", "sent", "assinatura enviada")
    users_path = Path(__file__).resolve().parent.parent / "config" / "users.json"
    app = create_app(test_config={
        "TESTING": True,
        "SECRET_KEY": "test-secret",
        "SHARED_SQLITE_PATH": str(database_path),
        "ADMIN_USERS_PATH": str(users_path),
        "REQUEST_LOG_PATH": str(log_path),
    })

    with app.test_client() as client:
        login_page = client.get("/admin/login")
        token = login_page.text.split('name="csrf_token" value="', 1)[1].split('"', 1)[0]
        client.post("/admin/login", data={"username": "rh", "password": "rh", "csrf_token": token})
        dashboard = client.get("/admin")
        assert dashboard.status_code == 200
        assert 'data-column="0"' in dashboard.text
        assert "Assinatura gerada" in dashboard.text
        assert 'class="signature-state generated">Sim' in dashboard.text
        assert "Cobertura Lorac" in dashboard.text
        assert "1 de 2" in dashboard.text
        assert "Sem E-mail Válido" in dashboard.text


def test_merge_emails_matches_normalized_names_without_overwriting(tmp_path):
    from openpyxl import Workbook

    employees_path = tmp_path / "colaboradores.xlsx"
    employees = Workbook()
    employees.active.append(["Nome", "Cargo", "Email"])
    employees.active.append(["Ana Luíza de Carvalho", "Analista", ""])
    employees.active.append(["E-mail Existente", "Gestor", "manter@empresa.com"])
    employees.save(employees_path)

    emails_path = tmp_path / "emails.xlsx"
    emails = Workbook()
    emails.active.append(["NOME", "Email"])
    emails.active.append(["  ANA LUIZA DE CARVALHO ", "ana.luiza@empresa.com"])
    emails.active.append(["E-MAIL EXISTENTE", "substituir@empresa.com"])
    emails.save(emails_path)

    result = merge_emails_from_workbook(employees_path, emails_path)

    assert result["updated"] == 1
    assert find_employee_by_email(employees_path, "ana.luiza@empresa.com")["full_name"] == "Ana Luíza de Carvalho"
    assert find_employee_by_email(employees_path, "manter@empresa.com")["full_name"] == "E-mail Existente"
    assert list(tmp_path.joinpath("backups").glob("*.xlsx"))
