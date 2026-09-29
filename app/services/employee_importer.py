import csv
import io
import re
from pathlib import Path

from openpyxl import load_workbook

EMAIL_PATTERN = re.compile(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")


def normalize_email(value):
    return (value or "").strip().lower()


def normalize_text(value):
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def normalize_name(value):
    return normalize_text(value)


def validate_employee_row(row, require_email=True):
    issues = {}
    full_name = normalize_name(row.get("nome") or row.get("full_name") or row.get("name"))
    job_title = normalize_text(row.get("cargo") or row.get("job_title") or row.get("role"))
    email = normalize_email(row.get("email") or row.get("e_mail") or row.get("e-mail"))

    if not full_name:
        issues["nome"] = "Nome obrigatório."
    if not job_title:
        issues["cargo"] = "Cargo obrigatório."
    if require_email and not email:
        issues["email"] = "E-mail obrigatório."
    elif email and not EMAIL_PATTERN.fullmatch(email):
        issues["email"] = "E-mail inválido."

    return issues


def read_employee_rows(file_obj, filename):
    file_obj.seek(0)
    filename_lower = (filename or "").lower()

    if filename_lower.endswith(".csv"):
        text = file_obj.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        return [
            {
                "nome": row.get("nome") or row.get("Nome") or row.get("name") or "",
                "cargo": row.get("cargo") or row.get("Cargo") or row.get("role") or "",
                "email": row.get("email") or row.get("Email") or row.get("e-mail") or "",
                "ativo": row.get("ativo") or row.get("Ativo") or row.get("active") or "SIM",
                "telefone": row.get("telefone") or row.get("Telefone") or row.get("celular") or row.get("Celular") or row.get("phone") or "",
            }
            for row in reader
        ]

    if filename_lower.endswith((".xlsx", ".xlsm")):
        workbook = load_workbook(file_obj, read_only=True)
        sheet = workbook.active
        rows = []
        headers = []
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if index == 1:
                headers = [str(cell or "").strip().lower() for cell in row]
                continue
            values = {headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))}
            rows.append({
                "nome": values.get("nome") or values.get("name") or values.get("full_name") or "",
                "cargo": values.get("cargo") or values.get("role") or values.get("job_title") or "",
                "email": values.get("email") or values.get("e-mail") or values.get("e_mail") or "",
                "ativo": values.get("ativo") or values.get("active") or "SIM",
                "telefone": values.get("telefone") or values.get("celular") or values.get("telefone celular") or values.get("celular corporativo") or values.get("phone") or "",
            })
        return rows

    raise ValueError("Formato de arquivo não suportado. Use CSV ou XLSX.")


def safe_filename(value):
    safe = re.sub(r"[^a-zA-Z0-9_\-]+", "_", value.strip())
    return safe.strip("_") or "assinatura"


def build_employee_record(row):
    return {
        "full_name": normalize_name(row.get("nome") or row.get("full_name") or row.get("name") or ""),
        "job_title": normalize_text(row.get("cargo") or row.get("job_title") or row.get("role") or ""),
        "email": normalize_email(row.get("email") or row.get("e_mail") or row.get("e-mail") or ""),
        "active": str(row.get("ativo") or row.get("active") or "SIM").strip().upper() in {"SIM", "YES", "Y"},
        "phone": normalize_text(row.get("telefone") or row.get("celular") or row.get("telefone celular") or row.get("celular corporativo") or row.get("phone") or ""),
    }
