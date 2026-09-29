import os
import shutil
import tempfile
import threading
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook


_WORKBOOK_LOCK = threading.RLock()


def _key(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(character for character in text if not unicodedata.combining(character)).strip().lower()


def _safe_excel_text(value):
    text = str(value or "").strip()
    return "'" + text if text.startswith("=") else text


def _is_active(value):
    return _key(value or "sim") not in {"nao", "no", "n", "0", "false", "inativo"}


def _headers(sheet):
    return {_key(cell.value): index for index, cell in enumerate(sheet[1], start=1) if cell.value is not None}


def _backup(path):
    backup_dir = path.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    target = backup_dir / f"{path.stem}_{stamp}{path.suffix}"
    shutil.copy2(path, target)
    return target


def _atomic_save(workbook, path):
    descriptor, temporary_name = tempfile.mkstemp(prefix=f".{path.stem}_", suffix=path.suffix, dir=path.parent)
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def ensure_employee_columns(workbook_path):
    """Create required columns, with backup and atomic replacement."""
    path = Path(workbook_path)
    if not path.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {path}")
    with _WORKBOOK_LOCK:
        workbook = load_workbook(path)
        try:
            sheet = workbook.active
            existing = _headers(sheet)
            changed = False
            for normalized, label in (("email", "Email"), ("celular", "Celular"), ("ativo", "Ativo")):
                if normalized not in existing:
                    sheet.cell(1, sheet.max_column + 1, label)
                    changed = True
            if changed:
                _backup(path)
                _atomic_save(workbook, path)
            return changed
        finally:
            workbook.close()


def list_employees(workbook_path):
    """Read every named employee directly from the official workbook."""
    path = Path(workbook_path)
    if not path.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {path}")
    with _WORKBOOK_LOCK:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            header_values = next(rows, ())
            headers = {_key(value): index for index, value in enumerate(header_values) if value is not None}
            employees = []
            for row_number, row in enumerate(rows, start=2):
                def value(*names):
                    for name in names:
                        column = headers.get(name)
                        if column is not None and column < len(row):
                            result = row[column]
                            if result is not None:
                                return str(result).strip()
                    return ""

                full_name = value("nome", "name", "full_name")
                if not full_name:
                    continue
                employees.append({
                    "id": row_number,
                    "full_name": full_name,
                    "job_title": value("cargo", "cargo oficial", "role", "job_title"),
                    "email": value("email", "e-mail", "e_mail").lower(),
                    "phone": value("celular", "telefone", "telefone celular", "celular corporativo", "phone"),
                    "active": _is_active(value("ativo", "active")),
                })
            return employees
        finally:
            workbook.close()


def find_employee_by_email(workbook_path, email, active_only=True):
    normalized = str(email or "").strip().lower()
    for employee in list_employees(workbook_path):
        if employee["email"] == normalized and (employee["active"] or not active_only):
            return employee
    return None


def get_employee(workbook_path, row_number):
    return next((item for item in list_employees(workbook_path) if item["id"] == row_number), None)


def save_employee_to_workbook(workbook_path, employee, original=None):
    """Update or append an employee, backing up and atomically replacing the workbook."""
    path = Path(workbook_path)
    if not path.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {path}")
    with _WORKBOOK_LOCK:
        workbook = load_workbook(path)
        try:
            sheet = workbook.active
            headers = _headers(sheet)
            required = {"nome": "Nome", "cargo": "Cargo", "email": "Email", "celular": "Celular", "ativo": "Ativo"}
            for normalized, label in required.items():
                if normalized not in headers:
                    column = sheet.max_column + 1
                    sheet.cell(1, column, label)
                    headers[normalized] = column

            original = original or {}
            row_number = original.get("id")
            if not row_number:
                original_name = _key(original.get("full_name"))
                original_email = _key(original.get("email"))
                for row in range(2, sheet.max_row + 1):
                    row_name = _key(sheet.cell(row, headers["nome"]).value)
                    row_email = _key(sheet.cell(row, headers["email"]).value)
                    if (original_email and row_email == original_email) or (original_name and row_name == original_name):
                        row_number = row
                        break
            if not row_number:
                row_number = sheet.max_row + 1

            values = {
                "nome": employee["full_name"], "cargo": employee["job_title"],
                "email": employee["email"], "celular": employee.get("phone", ""),
                "ativo": "SIM" if employee.get("active", True) else "NÃO",
            }
            for field, value in values.items():
                sheet.cell(int(row_number), headers[field], _safe_excel_text(value))

            _backup(path)
            _atomic_save(workbook, path)
            return int(row_number)
        finally:
            workbook.close()


def replace_employee_workbook(workbook_path, employees):
    """Replace the official workbook from a validated import, retaining a backup."""
    path = Path(workbook_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with _WORKBOOK_LOCK:
        workbook = Workbook()
        try:
            sheet = workbook.active
            sheet.title = "Colaboradores"
            sheet.append(["Nome", "Cargo", "Email", "Celular", "Ativo"])
            for employee in employees:
                sheet.append([
                    _safe_excel_text(employee["full_name"]), _safe_excel_text(employee["job_title"]),
                    _safe_excel_text(employee["email"]), _safe_excel_text(employee.get("phone", "")),
                    "SIM" if employee.get("active", True) else "NÃO",
                ])
            if path.exists():
                _backup(path)
            _atomic_save(workbook, path)
        finally:
            workbook.close()


def replace_employee_workbook_from_upload(workbook_path, file_obj):
    """Replace the official workbook with a validated uploaded XLSX in full."""
    path = Path(workbook_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with _WORKBOOK_LOCK:
        descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{path.stem}_upload_", suffix=path.suffix, dir=path.parent
        )
        os.close(descriptor)
        temporary_path = Path(temporary_name)
        try:
            file_obj.seek(0)
            with temporary_path.open("wb") as destination:
                shutil.copyfileobj(file_obj, destination)

            # Refuse a corrupt or non-Excel payload before touching production data.
            uploaded = load_workbook(temporary_path, read_only=True, data_only=False)
            uploaded.close()

            if path.exists():
                _backup(path)
            os.replace(temporary_path, path)
        finally:
            if temporary_path.exists():
                temporary_path.unlink()


def merge_emails_from_workbook(workbook_path, email_workbook_path):
    """Fill blank employee e-mails from unique exact normalized-name matches."""
    path = Path(workbook_path)
    email_path = Path(email_workbook_path)
    if not path.exists() or not email_path.exists():
        raise FileNotFoundError("A planilha de colaboradores e a planilha de e-mails são obrigatórias.")

    candidates = {}
    source = load_workbook(email_path, read_only=True, data_only=True)
    try:
        sheet = source.active
        rows = sheet.iter_rows(values_only=True)
        header = {_key(value): index for index, value in enumerate(next(rows, ())) if value is not None}
        name_column = header.get("nome")
        email_column = header.get("email")
        if name_column is None or email_column is None:
            raise ValueError("A planilha auxiliar precisa ter as colunas Nome e Email.")
        for row in rows:
            if name_column >= len(row) or email_column >= len(row):
                continue
            name = _key(row[name_column])
            email = str(row[email_column] or "").strip().lower()
            if name and email:
                candidates.setdefault(name, set()).add(email)
    finally:
        source.close()

    unique_emails = {name: next(iter(emails)) for name, emails in candidates.items() if len(emails) == 1}
    ambiguous_names = {name for name, emails in candidates.items() if len(emails) > 1}
    updated = 0
    unmatched = 0
    with _WORKBOOK_LOCK:
        workbook = load_workbook(path)
        try:
            sheet = workbook.active
            headers = _headers(sheet)
            name_column = headers.get("nome")
            if name_column is None:
                raise ValueError("A planilha oficial precisa ter a coluna Nome.")
            email_column = headers.get("email")
            if email_column is None:
                email_column = sheet.max_column + 1
                sheet.cell(1, email_column, "Email")
            for row_number in range(2, sheet.max_row + 1):
                if str(sheet.cell(row_number, email_column).value or "").strip():
                    continue
                name = _key(sheet.cell(row_number, name_column).value)
                email = unique_emails.get(name)
                if email:
                    sheet.cell(row_number, email_column, _safe_excel_text(email))
                    updated += 1
                elif name:
                    unmatched += 1
            if updated:
                _backup(path)
                _atomic_save(workbook, path)
        finally:
            workbook.close()
    return {"updated": updated, "unmatched": unmatched, "ambiguous": len(ambiguous_names)}
